import os
# import argparse
from docx import Document
from fpdf import FPDF


class DocxToPdfConverter:
    def __init__(self, font_style="Arial", font_size=12, line_height=6):
        self.font_style = font_style
        self.font_size = font_size
        self.line_height = line_height

    
    def get_docx_inputs(root_folder):
        inputs = []

        # Recorre todas las subcarpetas
        for subdir, _, files in os.walk(root_folder):
            for file in files:
                # Verifica si es un .docx que empieza con "0"
                if file.lower().endswith(".docx") and file.startswith("0"):
                    full_path = os.path.join(subdir, file)
                    inputs.append(full_path)

        return inputs
    
    def extract_text_from_docx(self, docx_path):
        """Extraer texto de un archivo DOCX"""
        try:
            doc = Document(docx_path)
            full_text = []
            for paragraph in doc.paragraphs:
                full_text.append(paragraph.text)
            return '\n'.join(full_text)
        except Exception as e:
            print(f"Error al leer {docx_path}: {str(e)}")
            return ""
    
    def convert_to_pdf(self, docx_files, output_pdf):
        """Convertir los documentos DOCX a un único PDF"""
        if not docx_files or not output_pdf:
            print("Error: Debe especificar archivos DOCX y un PDF de salida")
            return False
        
        try:
            # Crear PDF
            pdf = FPDF()
            pdf.set_auto_page_break(auto=True, margin=15)
            pdf.add_page()

            # Registrar y configurar fuente Verdana
            pdf.add_font("Verdana", "", "fonts/verdana.ttf", uni=True)
            pdf.add_font("Verdana", "B", "fonts/verdanab.ttf", uni=True)
            # Configurar fuente
            pdf.set_font(self.font_style, "", self.font_size)
            
            print(f"Procesando {len(docx_files)} documentos...")
            
            # Extraer y agregar texto de cada documento
            for i, docx_file in enumerate(docx_files, 1):
                print(f"Procesando documento {i}: {os.path.basename(docx_file)}")
                text = self.extract_text_from_docx(docx_file)
                
                # Agregar número de documento
                pdf.set_font(self.font_style, "B", self.font_size + 2)
                pdf.cell(0, 10, f"Documento #{i}: {os.path.basename(docx_file)}", ln=True)
                pdf.set_font(self.font_style, "", self.font_size)
                
                # Agregar contenido
                pdf.multi_cell(0, self.line_height, text)
                pdf.ln(5)
            
            # Guardar PDF
            pdf.output(output_pdf)
            print(f"PDF guardado exitosamente en: {output_pdf}")
            return True
        
        except Exception as e:
            print(f"Error durante la conversión: {str(e)}")
            return False

def main():
#     parser = argparse.ArgumentParser(description='Convertir documentos DOCX a PDF')
#     parser.add_argument('--input', '-i', nargs='+', required=True, 
#                        help='Archivos DOCX de entrada (separados por espacios)')
#     parser.add_argument('--output', '-o', required=True, 
#                        help='Archivo PDF de salida')
#     parser.add_argument('--font', '-f', default='Arial', 
#                        help='Estilo de fuente (por defecto: Arial)')
#     parser.add_argument('--size', '-s', type=int, default=12, 
#                        help='Tamaño de fuente (por defecto: 12)')
    
#     args = parser.parse_args()
    
#     # Verificar que los archivos de entrada existen
#     for file_path in args.input:
#         if not os.path.exists(file_path):
#             print(f"Error: El archivo {file_path} no existe")
#             return

    carpeta_raiz = "C:\\Users\\Orlando\\Desktop\\Py-DocxExtractor"
    docx_inputs = DocxToPdfConverter.get_docx_inputs(carpeta_raiz)

    print("Archivos encontrados:")

    input = []
    for i, docx in enumerate(docx_inputs, 1):
        print(f"{i}. {docx}")
        input.append(docx)

    
    modes = [
        {
            "font_style": "arial",
            "font_size": 28,
            "line_height": 12 * 1.5,
            "output": "Arial_28_1-5.pdf"
        },
        {
            "font_style": "verdana",
            "font_size": 32,
            "line_height": 12 * 1.5,
            "output": "Verdana_32_1-5.pdf"
        },
        {
            "font_style": "verdana",
            "font_size": 32,
            "line_height": 12 * 1.15,
            "output": "Verdana_32_1_15.pdf"
        },
        {
            "font_style": "verdana",
            "font_size": 14,
            "line_height": 9 * 1.15,
            "output": "Verdana_14_1_15.pdf"
        }
    ] 

    for mode in modes:
        # Crear y ejecutar el conversor
        print(f"Fuente: {mode['font_style']} - Tamaño: {mode['font_size']} - LineHeight: {mode['line_height']} - PDF: {mode['output']}")
        # converter = DocxToPdfConverter(font_style=args.font, font_size=args.size)
        converter = DocxToPdfConverter(font_style=mode['font_style'], font_size=mode['font_size'], line_height = mode['line_height'])
        # converter.convert_to_pdf(args.input, args.output) 
        converter.convert_to_pdf(input, mode['output'])



if __name__ == "__main__":
    main()


